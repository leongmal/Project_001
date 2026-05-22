### Проверка : если в базе "Опрос ЭСК" отсутствуют номера
### счетчиков из отчета С.Кривцуна ,  будет выдан список несовпадений


from openpyxl import load_workbook
import openpyxl
AS.xlsx
#file = open(r'C:\Users\Leonik\Desktop\архивная\No_AS.txt','w')
wb =load_workbook(r'C:\Users\Leonik\Desktop\архивная\AS.xlsx')

names = []
wsmx = wb.sheetnames
for st in wsmx:
    ws = wb[st]
    for row in ws:
        name = row[3].value
        if name == None :
            if row[1].value == None:
                continue
            if (str(row[1].value)).isdigit() != True :
                continue
            else:
                names.append(int(row[1].value))
#print(names)
wbs =load_workbook(r'C:\Users\Leonik\Desktop\Опрос ЭСК от 02.12.2024.xlsx')
wss = wbs.active
z=[]
n = 0
for row in wss:
    name1 = row[7].value
    
    z.append(name1)
#print(z)
q=[]
for i in names:
    if i not in z:
        q.append(i)
print(q)
    
