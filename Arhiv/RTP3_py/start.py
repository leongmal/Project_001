import pandas as pd
from openpyxl import load_workbook
import json
from pathlib import Path


with open('config.json', 'r', encoding='utf-8') as f:
    config = json.load(f)
# Параметры
#excel_file=Path(r'\\Esk-node2\ора\РАСПРЕДЕЛЕНИЕ ЭЭ\2026\01. Январь\Распределение ЭЭ январь 2026.xlsx')
excel_file = Path(config['excel_file'])
column_numbers = '№ счётчика'
column_values = 'Расход, кВтч'
search_list_file = '/Users/Leonik/AppData/Local/Programs/Python/Python312/Project_001/RTP3/serial_numbers.txt'

def read_search_list_from_file(file_path):
    try:
        with open(file_path, 'r', encoding='UTF-8') as file:
            # Предполагаем, что в файле каждое число на новой строке
            search_list = [line.strip() for line in file if line.strip()]
            # Преобразуем строки в числа
            search_list = [int(num) for num in search_list]
        return search_list
    except Exception as e:
        print(f"Ошибка при чтении файла search_list: {e}")
        return []

def create_number_dict(excel_file, column_numbers, column_values, search_list):
    # Читаем Excel-файл
    df = pd.read_excel(excel_file, sheet_name='РП')
    result_dict = {}

    # Проверяем наличие столбцов
    if column_numbers not in df.columns:
        raise ValueError(f"Столбец '{column_numbers}' не найден в файле.")
    if column_values not in df.columns:
        raise ValueError(f"Столбец '{column_values}' не найден в файле.")

    # Преобразуем search_list в строки для корректного сравнения
    search_list_str = list(map(str, search_list))

    # Итерируем по строкам DataFrame
    for index, row in df.iterrows():
        current_number = row[column_numbers]

        # Пропускаем пропущенные значения (NaN, None и т.п.)
        if pd.isna(current_number):
            continue

        # Приводим к строке и убираем пробелы по краям
        current_number = str(current_number).strip()

        # Пропускаем значения, которые превратились в 'nan' после str()
        if current_number.lower() == 'nan':
            continue

        # Проверяем, есть ли номер в списке поиска (сравнение строк)
        if current_number in search_list_str:
            value = row[column_values]

            # Если значение расхода не пропущено — добавляем в словарь
            if pd.notna(value):
                result_dict[current_number] = int(value)
            else:
                # Если расход пропущен — записываем None
                result_dict[current_number] = None

    return result_dict

# Читаем search_list из файла
search_list = read_search_list_from_file(search_list_file)

try:
    result = create_number_dict(excel_file, column_numbers, column_values, search_list)
    
    print("Результат:")
    #print(result)
    
    with open('/Users/Leonik/AppData/Local/Programs/Python/Python312/Project_001/RTP3/num_val.txt', 'w', encoding='UTF-8') as file:
        json.dump(result, file, ensure_ascii=False, indent=4)
        
except Exception as e:
    print(f"Ошибка: {e}")


#file_path = r'C:\Users\Leonik\Documents\Ввод замеров по средним нагрузкам в сети 6-220 кВ - Январь 2026 г..xlsx'
file_path =Path(config["file_path"])
sheet_name = 'Разомкнутые сети-ТП'

# Путь к файлу с данными
result_file_path = '/Users/Leonik/AppData/Local/Programs/Python/Python312/Project_001/RTP3/num_val.txt'

def load_result_dict_from_file(file_path):
    try:
        with open(file_path, 'r', encoding='UTF-8') as file:
            result_dict = json.load(file)
        return result_dict
    except Exception as e:
        print(f"Ошибка при чтении файла с данными: {e}")
        return {}

# Загружаем словарь result из файла
result = load_result_dict_from_file(result_file_path)

# Открываем файл Excel
try:
    wb = load_workbook(file_path)
    ws = wb[sheet_name]
    
    # Проходим по строкам листа (начиная с 1-й)
    for row in range(1, ws.max_row + 1):
        # Читаем значение из 3-го столбца (столбец C = индекс 3)
        cell_value = ws.cell(row=row, column=3).value
        
        # Если значение есть и не пустое
        if cell_value is not None:
            # Приводим к строке и убираем пробелы
            key = str(cell_value).strip()
            
            # Если ключ есть в словаре result — записываем значение в 8-й столбец (H)
            if key in result:
                ws.cell(row=row, column=8).value = result[key] /1000 # столбец H = индекс 8
                ws.cell(row=row, column=10).value = 0.75 # столбец 10 = cos
                ws.cell(row=row, column=13).value = 0.5 # столбец 13 = kзап, о.е.

    # Сохраняем изменения в исходный файл
    wb.save(file_path)
    
    print('\n',"===== Изменения успешно сохранены в файл. =====",'\n')
    
except Exception as e:
    print(f"Произошла ошибка: {e}")
