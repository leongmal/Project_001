
from openpyxl import load_workbook
import json

# Путь к файлу Excel (должен быть .xlsx!)
file_path = r'C:\Users\Leonik\Documents\Ввод замеров по средним нагрузкам в сети 6-220 кВ - Январь 2026 г..xlsx'
sheet_name = 'Разомкнутые сети-ТП'

# Путь к файлу с данными
result_file_path = '/Users/Leonik/AppData/Local/Programs/Python/Python312/RTP3_py/num_val.txt'

def load_result_dict_from_file(file_path):
    try:
        with open(file_path, 'r', encoding='UTF-8') as file:
            result_dict = json.load(file)
        return result_dict
    except Exception as e:
        print(f"Ошибка при чтении файла с данными: {e}")
        return {}

# Загружаем словарь result из файла
result = load_result_dict_from_file(result_file_path)

# Открываем файл Excel
try:
    wb = load_workbook(file_path)
    ws = wb[sheet_name]
    
    # Проходим по строкам листа (начиная с 1-й)
    for row in range(1, ws.max_row + 1):
        # Читаем значение из 3-го столбца (столбец C = индекс 3)
        cell_value = ws.cell(row=row, column=3).value
        
        # Если значение есть и не пустое
        if cell_value is not None:
            # Приводим к строке и убираем пробелы
            key = str(cell_value).strip()
            
            # Если ключ есть в словаре result — записываем значение в 8-й столбец (H)
            if key in result:
                ws.cell(row=row, column=8).value = result[key] /1000 # столбец H = индекс 8

    # Сохраняем изменения в исходный файл
    wb.save(file_path)
    print("Изменения успешно сохранены в файл.")
    
except Exception as e:
    print(f"Произошла ошибка: {e}")
