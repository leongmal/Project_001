import json
from openpyxl import load_workbook

INPUT_FILE = "data.xlsx"
OUTPUT_FILE = "companies.jsonl"
SHEET_NAME = "Sheet1"  # имя листа

wb = load_workbook(filename=INPUT_FILE, data_only=True)
ws = wb[SHEET_NAME]

# Получаем заголовки из первой строки
headers = [cell.value for cell in ws[1]]

with open(OUTPUT_FILE, "w", encoding="utf-8") as outfile:
    # Начинаем со второй строки (данные)
    for row in ws.iter_rows(min_row=2, values_only=True):
        # Пропускаем пустые строки
        if not any(row):
            continue

        # Собираем словарь по заголовкам
        data = {headers[i]: row[i] for i in range(len(headers)) if headers[i] is not None}

        # Записываем одну JSON-строку
        outfile.write(json.dumps(data, ensure_ascii=False) + "\n")

print(f"Готово: {OUTPUT_FILE}")
