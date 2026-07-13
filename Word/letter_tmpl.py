from docx import Document
import re
import json
import os

# --- Твоя функция замены (без изменений) ---
def replace_placeholders_preserving_format(paragraph, data):
    full_text = ''.join(run.text for run in paragraph.runs)
    pattern = re.compile(r'\{(\w+)\}')
    parts = []
    last_end = 0

    for match in pattern.finditer(full_text):
        start, end = match.span()
        if start > last_end:
            parts.append(('text', full_text[last_end:start]))
        parts.append(('placeholder', match.group(1)))
        last_end = end
    if last_end < len(full_text):
        parts.append(('text', full_text[last_end:]))

    new_text = ''.join(data.get(p, '{' + p + '}') if t == 'placeholder' else p for t, p in parts)

    first_run_style = None
    if paragraph.runs:
        fr = paragraph.runs[0]
        first_run_style = {
            'name': fr.font.name,
            'size': fr.font.size,
            'bold': fr.font.bold,
            'italic': fr.font.italic,
            'underline': fr.font.underline,
            'color': fr.font.color.rgb,
        }

    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)

    if new_text:
        run = paragraph.add_run(new_text)
        if first_run_style:
            run.font.name = first_run_style['name']
            run.font.size = first_run_style['size']
            run.font.bold = first_run_style['bold']
            run.font.italic = first_run_style['italic']
            run.font.underline = first_run_style['underline']
            run.font.color.rgb = first_run_style['color']


# --- Функция генерации одного письма ---
def create_letter(data: dict, template_path: str = "letter_template.docx", output_dir: str = "result") -> str:
    doc = Document(template_path)

    # Обработка таблиц
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_placeholders_preserving_format(paragraph, data)

    # Обработка основного текста
    for paragraph in doc.paragraphs:
        replace_placeholders_preserving_format(paragraph, data)

    consumer_clean = ''.join(c for c in data.get('consumer', 'unnamed') if c.isalnum() or c in ' -')
    filename = f"письмо_{consumer_clean}.docx"
    output_path = os.path.join(output_dir, filename)

    os.makedirs(output_dir, exist_ok=True)
    doc.save(output_path)
    print("Письмо сохранено:", output_path)
    return output_path


# --- Главная логика: построчное чтение JSONL ---
if __name__ == '__main__':
    CONFIG_FILE = "companies.jsonl"
    TEMPLATE_FILE = "letter_template.docx"

    try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            for line_num, line in enumerate(f, start=1):
                line = line.strip()
                if not line:
                    continue

                try:
                    data = json.loads(line)  # Каждая строка — отдельный словарь
                    if not isinstance(data, dict):
                        print(f"Строка {line_num}: не словарь, пропущена.")
                        continue
                    create_letter(data, template_path=TEMPLATE_FILE)
                except json.JSONDecodeError as e:
                    print(f"Ошибка в строке {line_num} (невалидный JSON): {e}")
                    # Не прерываем весь процесс — идем дальше
    except FileNotFoundError:
        print(f"Файл {CONFIG_FILE} не найден! Создайте его и добавьте туда строки в формате JSON.")
